"""
Unit and integration tests for RollCall AI Streamlit UI logic (Prompt 7).

Tests:
1. File upload processing (clean CSV vs invalid CSV).
2. Table corrections and re-validation confirmation flow.
3. Photo upload confidence threshold handling in UI flow.
4. Attendance metric calculations (4-state model, no zero division, 66.67% rate).
5. Chat interaction with AttendanceAgent and tool logging.
6. Export generation (Excel and PDF buffers).
7. Database commit with transaction integrity.
"""

import io
import sqlite3
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from app.agent import AttendanceAgent
from app.db import init_db, insert_student, insert_room, insert_session, get_attendance_by_session
from app.ui import (
    process_uploaded_file,
    revalidate_and_summarise,
    calculate_metrics_summary,
    commit_session_to_database,
    generate_export_buffers,
)
from tools.summary_tool import SummaryReport


class MockUploadedFile:
    """Mock Streamlit UploadedFile object."""

    def __init__(self, name: str, data: bytes):
        self.name = name
        self._data = data

    def getbuffer(self) -> bytes:
        return self._data

    def read(self) -> bytes:
        return self._data


@pytest.fixture
def test_db():
    conn = sqlite3.connect(":memory:")
    init_db(conn)
    insert_student(conn, usn="1MS21CS001", name="Alice Smith", branch="CSE")
    insert_student(conn, usn="1MS21CS002", name="Bob Jones", branch="CSE")
    insert_student(conn, usn="1MS21IS001", name="Charlie Brown", branch="ISE")
    insert_room(conn, room_id="LH-101", name="Lecture Hall 101")
    insert_room(conn, room_id="LH-102", name="Lecture Hall 102")
    insert_session(conn, session_id="SESS-20261024-01", date="2026-10-24", subject="Computer Networks")
    yield conn
    conn.close()


def test_ui_upload_clean_csv(tmp_path):
    """Verify that a clean CSV upload completes validation and summary without confirmation."""
    csv_content = (
        "usn,name,branch,room_id,status\n"
        "1MS21CS001,Alice Smith,CSE,LH-101,PRESENT\n"
        "1MS21CS002,Bob Jones,CSE,LH-101,ABSENT\n"
    ).encode("utf-8")
    uploaded = MockUploadedFile("clean_test.csv", csv_content)

    agent = AttendanceAgent()

    resp, tmp_file = process_uploaded_file(
        agent,
        uploaded,
        master_usns=["1MS21CS001", "1MS21CS002"],
        master_rooms=["LH-101"],
    )
    assert resp.success is True
    assert resp.needs_confirmation is False
    assert resp.summary_report is not None
    assert resp.summary_report["overall"]["confirmed_percentage"] == 50.0
    assert len(resp.tool_calls) >= 3  # read_csv_tool, validate_tool, summary_tool
    if tmp_file.exists():
        tmp_file.unlink()


def test_ui_upload_invalid_csv_triggers_confirmation():
    """Verify that an invalid CSV (duplicate USN) triggers needs_confirmation=True."""
    csv_content = (
        "usn,name,branch,room_id,status\n"
        "1MS21CS001,Alice Smith,CSE,LH-101,PRESENT\n"
        "1MS21CS001,Alice Smith,CSE,LH-101,ABSENT\n"
    ).encode("utf-8")
    uploaded = MockUploadedFile("bad_test.csv", csv_content)

    agent = AttendanceAgent()

    resp, tmp_file = process_uploaded_file(
        agent,
        uploaded,
        master_usns=["1MS21CS001"],
        master_rooms=["LH-101"],
    )
    assert resp.success is False
    assert resp.needs_confirmation is True
    assert resp.confirmation_request is not None
    assert any("DUPLICATE_USN" in err for err in resp.confirmation_request.issues)
    # Pipeline invariant: summary_tool MUST NOT have run
    tool_names = [call["tool_name"] for call in resp.tool_calls]
    assert "summary_tool" not in tool_names
    if tmp_file.exists():
        tmp_file.unlink()


def test_ui_table_corrections_and_revalidation():
    """Verify that correcting records and calling revalidate_and_summarise clears confirmation."""
    corrected_records = [
        {"usn": "1MS21CS001", "name": "Alice Smith", "branch": "CSE", "room_id": "LH-101", "status": "PRESENT"},
        {"usn": "1MS21CS002", "name": "Bob Jones", "branch": "CSE", "room_id": "LH-101", "status": "ABSENT"},
    ]

    agent = AttendanceAgent()

    resp, tmp_file = revalidate_and_summarise(
        agent,
        corrected_records,
        master_usns=["1MS21CS001", "1MS21CS002"],
        master_rooms=["LH-101"],
    )
    assert resp.success is True
    assert resp.needs_confirmation is False
    assert resp.summary_report is not None
    assert resp.summary_report["overall"]["present"] == 1
    assert resp.summary_report["overall"]["absent"] == 1
    assert resp.summary_report["overall"]["confirmed_percentage"] == 50.0
    if tmp_file.exists():
        tmp_file.unlink()


def test_ui_photo_upload_low_confidence_requires_confirmation(monkeypatch):
    """Verify photo upload with confidence < 0.80 halts pipeline in UI helper."""
    from unittest.mock import MagicMock
    from tools.read_photo_tool import PhotoExtractionResult, PhotoRow, ExtractionSource

    mock_row = PhotoRow(
        usn="1MS21CS001",
        name="Alice Smith",
        status="PRESENT",
        confidence=0.75,
        source=ExtractionSource.OCR_FALLBACK,
    )
    mock_result = PhotoExtractionResult(
        image_path="test_sheet.jpg",
        rows=[mock_row],
        source=ExtractionSource.OCR_FALLBACK,
    )

    agent = AttendanceAgent()
    agent.registry.tools["read_photo_tool"].handler = MagicMock(return_value=mock_result)

    uploaded = MockUploadedFile("test_sheet.jpg", b"fake_image_bytes")
    resp, tmp_file = process_uploaded_file(agent, uploaded)

    assert resp.success is False
    assert resp.needs_confirmation is True
    assert "0.80" in resp.message
    tool_names = [call["tool_name"] for call in resp.tool_calls]
    assert "summary_tool" not in tool_names
    if tmp_file.exists():
        tmp_file.unlink()


def test_ui_attendance_calculations():
    """Verify calculation helper adheres to 4-state logic and handles edge cases cleanly."""
    # 1. Empty / None summary
    empty_res = calculate_metrics_summary(None)
    assert empty_res["confirmed_percentage"] == 0.0
    assert empty_res["total"] == 0

    # 2. 2 PRESENT, 1 ABSENT -> 66.67%
    rep_66 = {
        "overall": {
            "present": 2,
            "absent": 1,
            "pending": 0,
            "unreviewed": 0,
            "confirmed": 3,
            "total": 3,
            "confirmed_percentage": 66.67,
        }
    }
    res_66 = calculate_metrics_summary(rep_66)
    assert res_66["confirmed_percentage"] == 66.67
    assert res_66["present"] == 2
    assert res_66["absent"] == 1

    # 3. Only PENDING -> 0.0% without division-by-zero
    rep_pending = {
        "overall": {
            "present": 0,
            "absent": 0,
            "pending": 2,
            "unreviewed": 0,
            "confirmed": 0,
            "total": 2,
            "confirmed_percentage": 0.0,
        }
    }
    res_pending = calculate_metrics_summary(rep_pending)
    assert res_pending["confirmed_percentage"] == 0.0
    assert res_pending["pending"] == 2
    assert res_pending["absent"] == 0  # Must never silently convert to absent


def test_ui_chat_agent_query(test_db):
    """Verify chat tab handles agent database questions and anti-hallucination queries."""
    agent = AttendanceAgent(db_target=test_db)

    # 1. Valid question
    resp = agent.process_request("Who was absent in session SESS-20261024-01?")
    assert resp.request_id is not None
    assert any(c["tool_name"] == "query_tool" for c in resp.tool_calls)

    # 2. Metric question outside schema -> Anti-hallucination check
    unanswerable_resp = agent.process_request("What was the average exam score in session SESS-20261024-01?")
    assert "I do not have that data." in unanswerable_resp.message


def test_ui_export_generation():
    """Verify Excel and PDF export buffers are generated successfully and valid."""
    room_df = pd.DataFrame([
        {"room_id": "LH-101", "present": 2, "absent": 1, "pending": 0, "unreviewed": 0, "unresolved": 0, "confirmed": 3, "total": 3, "percentage": 66.67}
    ])
    branch_df = pd.DataFrame([
        {"branch": "CSE", "present": 2, "absent": 1, "pending": 0, "unreviewed": 0, "unresolved": 0, "confirmed": 3, "total": 3, "percentage": 66.67}
    ])
    overall = {"present": 2, "absent": 1, "pending": 0, "unreviewed": 0, "unresolved": 0, "confirmed": 3, "total": 3, "percentage": 66.67, "confirmed_percentage": 66.67}

    summary = SummaryReport(room_wise=room_df, branch_wise=branch_df, overall=overall)
    excel_bytes, pdf_bytes = generate_export_buffers(summary)

    # Excel validation
    assert len(excel_bytes) > 0
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Room_Summary" in wb.sheetnames
    assert "Branch_Summary" in wb.sheetnames

    # PDF validation
    assert len(pdf_bytes) > 0
    assert pdf_bytes.startswith(b"%PDF-")


def test_ui_database_commit(tmp_path):
    """Verify committing confirmed session attendance writes to SQLite."""
    db_file = tmp_path / "test_commit.db"
    conn = sqlite3.connect(db_file)
    init_db(conn)
    insert_student(conn, usn="1MS21CS001", name="Alice Smith", branch="CSE")
    insert_room(conn, room_id="LH-101", name="Lecture Hall 101")
    conn.close()

    records = [
        {"usn": "1MS21CS001", "name": "Alice Smith", "room_id": "LH-101", "status": "PRESENT", "confidence": 1.0}
    ]

    ok, msg = commit_session_to_database(
        db_target=db_file,
        session_id="SESS-NEW-01",
        records=records,
        date="2026-10-25",
        subject="AI Lab",
    )
    assert ok is True
    assert "Successfully committed" in msg

    from app.db import get_connection
    verify_conn = get_connection(db_file)
    saved = get_attendance_by_session(verify_conn, "SESS-NEW-01")
    verify_conn.close()
    assert len(saved) == 1
    assert saved[0]["usn"] == "1MS21CS001"
    assert saved[0]["status"] == "PRESENT"
